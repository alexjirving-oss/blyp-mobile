import json
import re
from pathlib import Path

from openpyxl import load_workbook


PROJECT = Path(r"C:\Users\Alex\Blyp26\_agent\admin-workbook-20260802\admin")
WORKBOOK = PROJECT / "public" / "BLYP_Authoritative_Product_Workbook.xlsx"
WEB_DATA = PROJECT / "src" / "data" / "workbook.ts"
OUTPUT = PROJECT / "WORKBOOK_PARITY.json"


def extract_json_assignment(source: str, name: str):
    pattern = rf"export const {re.escape(name)}(?:[^=]*)= (.*?);\n(?=export const|$)"
    match = re.search(pattern, source, re.DOTALL)
    if not match:
        raise ValueError(f"Unable to locate {name} in generated admin data.")
    return json.loads(match.group(1))


source = WEB_DATA.read_text(encoding="utf-8")
stats = extract_json_assignment(source, "workbookStats")
arrays = {
    "Area Summary": extract_json_assignment(source, "areaSummaries"),
    "Feature Register": extract_json_assignment(source, "featureRecords"),
    "Evidence Ledger": extract_json_assignment(source, "evidenceRecords"),
    "Runtime Paths": extract_json_assignment(source, "runtimePathRecords"),
    "Component Map": extract_json_assignment(source, "componentRecords"),
    "Production Gates": extract_json_assignment(source, "productionGateRecords"),
    "Source Register": extract_json_assignment(source, "sourceRecords"),
    "Decision Log": extract_json_assignment(source, "decisionRecords"),
}

workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
expected_stats = {
    "Feature Register": ("featureCount", 633),
    "Area Summary": ("areaCount", 30),
    "Evidence Ledger": ("evidenceCount", 312),
    "Runtime Paths": ("runtimePathCount", 18),
    "Component Map": ("componentCount", 20),
    "Production Gates": ("gateCount", 14),
    "Source Register": ("sourceCount", 11),
    "Decision Log": ("decisionCount", 5),
}
sheet_rules = {
    "Feature Register": ("Feature Register", 1),
    "Evidence Ledger": ("Evidence Ledger", 1),
    "Runtime Paths": ("Runtime Paths", 1),
    "Component Map": ("Component Map", 1),
    "Production Gates": ("Production Gates", 4),
    "Source Register": ("Sources & Revisions", 1),
    "Decision Log": ("Decisions & Changes", 1),
}

feature_sheet = workbook["Feature Register"]
feature_headers = [cell.value for cell in next(feature_sheet.iter_rows(min_row=1, max_row=1))]
product_area_column = feature_headers.index("Product Area") + 1
workbook_area_count = len({
    row[0]
    for row in feature_sheet.iter_rows(
        min_row=2,
        min_col=product_area_column,
        max_col=product_area_column,
        values_only=True,
    )
    if row[0]
})

checks = []
for sheet_name, (stat_name, expected_count) in expected_stats.items():
    if sheet_name == "Area Summary":
        workbook_count = workbook_area_count
    else:
        workbook_sheet_name, header_rows = sheet_rules[sheet_name]
        workbook_count = max(workbook[workbook_sheet_name].max_row - header_rows, 0)
    frontend_count = len(arrays[sheet_name])
    stat_count = stats.get(stat_name)
    passed = workbook_count == frontend_count == stat_count == expected_count
    checks.append({
        "name": f"{sheet_name} parity",
        "status": "passed" if passed else "failed",
        "expected": expected_count,
        "workbookRows": workbook_count,
        "adminRecords": frontend_count,
        "reportedStat": stat_count,
    })

feature_columns = feature_sheet.max_column
checks.append({
    "name": "Feature Register schema width",
    "status": "passed" if feature_columns == 64 else "failed",
    "expectedColumns": 64,
    "actualColumns": feature_columns,
})

id_arrays = {
    name: records
    for name, records in arrays.items()
    if records and isinstance(records[0], dict) and "id" in records[0]
}
id_sheet_rules = {
    "Feature Register": ("Feature Register", 2),
    "Evidence Ledger": ("Evidence Ledger", 2),
    "Runtime Paths": ("Runtime Paths", 2),
    "Component Map": ("Component Map", 2),
    "Source Register": ("Sources & Revisions", 2),
    "Decision Log": ("Decisions & Changes", 2),
}
for name, records in id_arrays.items():
    identifiers = [record["id"] for record in records]
    workbook_sheet_name, first_data_row = id_sheet_rules[name]
    workbook_identifiers = [
        row[0]
        for row in workbook[workbook_sheet_name].iter_rows(
            min_row=first_data_row,
            min_col=1,
            max_col=1,
            values_only=True,
        )
        if row[0]
    ]
    exact_match = identifiers == workbook_identifiers
    unique = len(identifiers) == len(set(identifiers))
    checks.append({
        "name": f"{name} identifier parity and uniqueness",
        "status": "passed" if exact_match and unique else "failed",
        "records": len(identifiers),
        "uniqueIdentifiers": len(set(identifiers)),
        "workbookIdentifiers": len(workbook_identifiers),
        "exactOrderMatch": exact_match,
    })

workbook_gate_names = [
    row[0]
    for row in workbook["Production Gates"].iter_rows(
        min_row=5,
        min_col=1,
        max_col=1,
        values_only=True,
    )
    if row[0]
]
admin_gate_names = [record["name"] for record in arrays["Production Gates"]]
checks.append({
    "name": "Production gate name and order parity",
    "status": "passed" if admin_gate_names == workbook_gate_names else "failed",
    "records": len(admin_gate_names),
    "exactOrderMatch": admin_gate_names == workbook_gate_names,
})

revision = "8bc0dcb2a1da1b21368c5536aaeef829d587f316"
checks.append({
    "name": "Audited revision lineage",
    "status": "passed" if revision in source else "failed",
    "revision": revision,
})

result = {
    "status": "passed" if all(check["status"] == "passed" for check in checks) else "failed",
    "workbook": str(WORKBOOK),
    "adminData": str(WEB_DATA),
    "sheetNames": workbook.sheetnames,
    "checks": checks,
}
OUTPUT.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
print(json.dumps(result, indent=2))

if result["status"] != "passed":
    raise SystemExit(1)
